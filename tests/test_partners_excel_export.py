"""Экспорт базы партнёров в Excel (0110)."""

import io

from openpyxl import load_workbook

from app.common.module_access import AccessLevel, Module
from app.partners.excel import COLUMNS


def _partner(**kwargs):
    base = dict(category="REALTOR", name="Пётр Риэлторов", city="Москва", phone="8 900 123-45-67")
    base.update(kwargs)
    return base


def _sheet(resp):
    assert resp.status_code == 200
    return load_workbook(io.BytesIO(resp.content)).active


def _rows(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows(min_row=2)]


def test_export_columns_and_content(api, make_user):
    client = api(make_user(Module.CLIENTS, level=AccessLevel.EDIT))
    partner = client.post(
        "/api/partners/",
        json=_partner(
            organization="Этажи",
            contacts=[{"messenger": "telegram", "contact": "@petr"}, {"messenger": "max", "contact": "8 900 123-45-67"}],
            comment="Встреча 10.10",
        ),
    ).json()
    client.post(f"/api/partners/{partner['id']}/notes", json={"text": "комиссия 3%"})
    client.post(f"/api/partners/{partner['id']}/notes", json={"text": "приведёт клиента в пятницу"})

    resp = client.get("/api/partners/export.xlsx")
    assert resp.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in resp.headers["content-disposition"]
    assert ".xlsx" in resp.headers["content-disposition"]

    ws = _sheet(resp)
    assert ws.title == "Партнёры"
    assert [c.value for c in ws[1]] == [title for title, _ in COLUMNS]
    assert ws.freeze_panes == "A2"

    (row,) = _rows(ws)
    assert row[0] == "Пётр Риэлторов"
    assert row[1] == "Риэлтор"
    assert row[2] == "Этажи"
    assert row[3] == "Москва"
    assert row[4] == "8 900 123-45-67"  # ведущая 8 сохранилась текстом
    assert row[6] == "telegram: @petr\nmax: 8 900 123-45-67"
    assert row[7] == "Встреча 10.10"
    # Заметки от новых к старым, каждая — строкой с датой.
    lines = row[8].split("\n")
    assert len(lines) == 2
    assert lines[0].endswith("— приведёт клиента в пятницу")
    assert lines[1].endswith("— комиссия 3%")
    assert row[9]  # дата добавления


def test_export_follows_filters(api, make_user):
    client = api(make_user(Module.CLIENTS, level=AccessLevel.EDIT))
    client.post("/api/partners/", json=_partner())
    client.post("/api/partners/", json=_partner(name="ООО Торговля", category="COMMERCE", city="Казань"))
    client.post("/api/partners/", json=_partner(name="Анна Землякова", category="REALTOR", city="москва"))

    def names(**params):
        return [r[0] for r in _rows(_sheet(client.get("/api/partners/export.xlsx", params=params)))]

    assert len(names()) == 3
    assert names(category="COMMERCE") == ["ООО Торговля"]
    assert sorted(names(category="REALTOR", city="МОСКВА")) == ["Анна Землякова", "Пётр Риэлторов"]
    assert names(search="Торговля") == ["ООО Торговля"]


def test_export_empty_has_only_headers(api, make_user):
    client = api(make_user(Module.CLIENTS, level=AccessLevel.VIEW))
    ws = _sheet(client.get("/api/partners/export.xlsx", params={"search": "нет такого"}))
    assert [c.value for c in ws[1]] == [title for title, _ in COLUMNS]
    assert _rows(ws) == []


def test_export_does_not_turn_text_into_formulas(api, make_user):
    client = api(make_user(Module.CLIENTS, level=AccessLevel.EDIT))
    client.post("/api/partners/", json=_partner(comment="=1+1", organization="@SUM(A1)"))

    resp = client.get("/api/partners/export.xlsx")
    ws = load_workbook(io.BytesIO(resp.content)).active
    comment, organization = ws["H2"], ws["C2"]
    assert comment.value == "=1+1" and comment.data_type == "s"
    assert organization.value == "@SUM(A1)" and organization.data_type == "s"


def test_export_requires_clients_access(api, make_user):
    outsider = api(make_user(Module.WAREHOUSE, level=AccessLevel.FULL))
    assert outsider.get("/api/partners/export.xlsx").status_code == 403

    viewer = api(make_user(Module.CLIENTS, level=AccessLevel.VIEW))
    assert viewer.get("/api/partners/export.xlsx").status_code == 200
