"""Импорт выписки из банк-клиента (задача 0081-e).

Поверх импорта 0011-k: выписка грузится на конкретный счёт, контрагент каждой
строки сопоставляется с единым справочником (0081-c), повторная загрузка того
же файла на тот же счёт дублей не плодит, статья подставляется из истории
контрагента.
"""

import pytest

from app.accounting.models import Counterparty, CounterpartyKind, MoneyMovement, MoneySubkind
from app.common.module_access import Module

HEAD = "Дата,Контрагент,ИНН контрагента,Назначение платежа,Сумма,НДС,№ документа,Вид\n"


@pytest.fixture
def acc_user(make_user):
    return make_user(Module.ACCOUNTING, admin=True)


def _csv(*rows: str) -> bytes:
    return (HEAD + "".join(r + "\n" for r in rows)).encode("utf-8")


def _accounts(api_client) -> list[dict]:
    return api_client.get("/api/accounting/accounts").json()


def _import(api_client, content: bytes, account_id: int):
    return api_client.post(
        "/api/accounting/money-movements/import",
        files={"file": ("vypiska.csv", content, "text/csv")},
        params={"account_id": account_id},
    )


def test_statement_lands_on_the_chosen_account(db, api, acc_user):
    api_client = api(acc_user)
    first, second = _accounts(api_client)

    result = _import(api_client, _csv("01.09.2026,ООО Ромашка,7701234567,оплата,150000,25000,125,"), second["id"])
    assert result.status_code == 200
    data = result.json()
    assert data["imported"] == 1
    assert data["account_id"] == second["id"]
    assert second["name"] in data["account_label"]

    mm = db.query(MoneyMovement).one()
    assert mm.account_id == second["id"]
    assert mm.status.value == "draft"
    # реестр первого счёта этой проводки не видит
    on_first = api_client.get(
        "/api/accounting/money-movements", params={"account_id": first["id"]}
    ).json()
    assert on_first == []


def test_import_without_account_is_rejected(db, api, acc_user):
    r = api(acc_user).post(
        "/api/accounting/money-movements/import",
        files={"file": ("v.csv", _csv("01.09.2026,А,,оплата,100,0,1,"), "text/csv")},
    )
    assert r.status_code == 422
    assert db.query(MoneyMovement).count() == 0


def test_repeated_upload_of_the_same_statement_creates_no_duplicates(db, api, acc_user):
    api_client = api(acc_user)
    account_id = _accounts(api_client)[0]["id"]
    body = _csv(
        "01.09.2026,ООО Ромашка,7701234567,оплата,150000,25000,125,",
        "03.09.2026,ИФНС,7707083893,НДС за квартал,-274000,0,126,налоги и сборы",
    )

    first = _import(api_client, body, account_id).json()
    assert first["imported"] == 2 and first["duplicates"] == 0

    second = _import(api_client, body, account_id).json()
    assert second["imported"] == 0
    assert second["duplicates"] == 2
    assert db.query(MoneyMovement).count() == 2


def test_the_same_statement_on_another_account_is_not_a_duplicate(db, api, acc_user):
    api_client = api(acc_user)
    first, other = _accounts(api_client)
    body = _csv("01.09.2026,ООО Ромашка,7701234567,оплата,150000,25000,125,")

    _import(api_client, body, first["id"])
    result = _import(api_client, body, other["id"]).json()

    assert result["imported"] == 1 and result["duplicates"] == 0
    assert db.query(MoneyMovement).count() == 2


def test_counterparties_are_created_once_and_matched_afterwards(db, api, acc_user):
    api_client = api(acc_user)
    account_id = _accounts(api_client)[0]["id"]

    first = _import(
        api_client, _csv("01.09.2026,ООО «Ромашка»,7701234567,оплата,150000,25000,125,"), account_id
    ).json()
    assert first["counterparties_created"] == 1
    assert first["counterparties_matched"] == 0

    # другой платёж того же контрагента: наименование написано иначе, ИНН тот же
    second = _import(
        api_client, _csv("05.09.2026,Ромашка Торг,7701234567,оплата,60000,10000,130,"), account_id
    ).json()
    assert second["counterparties_created"] == 0
    assert second["counterparties_matched"] == 1
    assert db.query(Counterparty).count() == 1

    counterparty = db.query(Counterparty).one()
    assert counterparty.kind is CounterpartyKind.OTHER
    payments = api_client.get(f"/api/accounting/counterparties/{counterparty.id}/payments").json()
    assert len(payments) == 2


def test_subkind_is_taken_from_the_settled_history_of_the_counterparty(db, api, acc_user):
    api_client = api(acc_user)
    account_id = _accounts(api_client)[0]["id"]

    # первый платёж арендодателю — статья из файла, проводим его
    first_ids = _import(
        api_client, _csv("01.08.2026,ООО Аренда,7712345678,за июль,-130000,0,50,аренда"), account_id
    ).json()["created_ids"]
    for mm_id in first_ids:
        api_client.post(f"/api/accounting/money-movements/{mm_id}/status", json={"to": "approved"})
        api_client.post(f"/api/accounting/money-movements/{mm_id}/status", json={"to": "posted"})

    # второй платёж тому же контрагенту приходит без статьи в файле
    result = _import(
        api_client, _csv("01.09.2026,ООО Аренда,7712345678,за август,-130000,0,51,"), account_id
    ).json()
    assert result["preliminary_subkind"] == 0

    latest = db.query(MoneyMovement).filter(MoneyMovement.external_number == "51").one()
    assert latest.subkind is MoneySubkind.RENT


def test_ambiguous_history_leaves_the_subkind_preliminary(db, api, acc_user):
    api_client = api(acc_user)
    account_id = _accounts(api_client)[0]["id"]

    mixed = _import(
        api_client,
        _csv(
            "01.08.2026,ООО Разное,7712345679,за июль,-1000,0,60,аренда",
            "02.08.2026,ООО Разное,7712345679,налог,-2000,0,61,налоги и сборы",
        ),
        account_id,
    ).json()["created_ids"]
    for mm_id in mixed:
        api_client.post(f"/api/accounting/money-movements/{mm_id}/status", json={"to": "approved"})
        api_client.post(f"/api/accounting/money-movements/{mm_id}/status", json={"to": "posted"})

    result = _import(
        api_client, _csv("01.09.2026,ООО Разное,7712345679,за август,-3000,0,62,"), account_id
    ).json()
    # история разнобойная — гадать не будем, остаётся предварительная статья
    assert result["preliminary_subkind"] == 1
    latest = db.query(MoneyMovement).filter(MoneyMovement.external_number == "62").one()
    assert latest.subkind is MoneySubkind.OTHER_EXPENSE


def test_template_has_the_inn_column(db, api, acc_user):
    import io

    from openpyxl import load_workbook

    r = api(acc_user).get("/api/accounting/money-movements/import/template")
    assert r.status_code == 200
    workbook = load_workbook(io.BytesIO(r.content))
    headers = [cell.value for cell in next(workbook.active.iter_rows(max_row=1))]
    assert "ИНН контрагента" in headers
