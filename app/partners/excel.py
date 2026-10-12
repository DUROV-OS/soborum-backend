"""Выгрузка базы партнёров в Excel (0110)."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.partners.models import Partner, PartnerCategory

CATEGORY_LABELS = {
    PartnerCategory.COMMERCE: "Коммерция",
    PartnerCategory.REAL_ESTATE_AGENCY: "Агентство недвижимости",
    PartnerCategory.REALTOR: "Риэлтор",
    PartnerCategory.LAND_SPECIALIST: "Специалист по земле",
}

# (заголовок, ширина колонки)
COLUMNS = [
    ("Партнёр", 28),
    ("Категория", 24),
    ("Организация", 26),
    ("Город", 18),
    ("Телефон", 18),
    ("Почта", 26),
    ("Мессенджеры", 28),
    ("Комментарий", 40),
    ("Договорённости", 50),
    ("Добавлен", 12),
]


def _date(value) -> str:
    return value.strftime("%d.%m.%Y") if value else ""


def _row(partner: Partner) -> list[str]:
    messengers = "\n".join(f"{c['messenger']}: {c['contact']}" for c in partner.contacts or [])
    notes = "\n".join(f"{_date(n.created_at)} — {n.text}" for n in partner.notes)
    return [
        partner.name,
        CATEGORY_LABELS.get(partner.category, partner.category.value),
        partner.organization or "",
        partner.city,
        partner.phone or "",
        partner.email or "",
        messengers,
        partner.comment or "",
        notes,
        _date(partner.created_at),
    ]


def build_partners_xlsx(partners: list[Partner]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Партнёры"
    ws.append([title for title, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for partner in partners:
        ws.append(_row(partner))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            # Пользовательский текст пишем строго как строку: иначе openpyxl
            # превратит значение, начинающееся с «=», в формулу.
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
