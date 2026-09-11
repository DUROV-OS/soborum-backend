"""Импорт платежей в «Бухгалтерию» таблицей (xlsx/csv) с ИИ-разметкой колонок.

По образцу `app/warehouse/price_import.py` (задачи 0011-g / 0011-h):
читаем таблицу → сопоставляем заголовки с полями проводки (Claude forced tool,
откат на словарь синонимов) → строим строки. Критичных колонок нет — отказ.
Некритичные (`subkind`, `payment_purpose`) могут быть пустыми: вид ИИ проставит
отдельным шагом, дозаполнение — задачей по кнопке.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone

from fastapi import HTTPException, UploadFile, status

from app.accounting.models import MoneyDirection, MoneySubkind
from app.core.config import settings

MAX_DATA_ROWS = 2000
SAMPLE_ROWS_FOR_AI = 6

# Некритичные поля — для отчёта «чего не хватило».
OPTIONAL_FIELDS = ("subkind", "payment_purpose")

# Русские подписи подвидов -> enum (для колонки «вид», если она есть в файле).
_SUBKIND_BY_LABEL: dict[str, MoneySubkind] = {
    "доход от продажи": MoneySubkind.SALE_INCOME,
    "продажа": MoneySubkind.SALE_INCOME,
    "выплата зарплаты": MoneySubkind.SALARY_PAYOUT,
    "зарплата": MoneySubkind.SALARY_PAYOUT,
    "оплата поставки": MoneySubkind.SUPPLY_PAYMENT,
    "закупка товаров": MoneySubkind.SUPPLY_PAYMENT,
    "поставка": MoneySubkind.SUPPLY_PAYMENT,
    "налоги и сборы": MoneySubkind.TAX,
    "налог": MoneySubkind.TAX,
    "ндс": MoneySubkind.TAX,
    "аренда": MoneySubkind.RENT,
    "прочий доход": MoneySubkind.OTHER_INCOME,
    "прочий расход": MoneySubkind.OTHER_EXPENSE,
}


def subkind_from_text(value: str | None) -> MoneySubkind | None:
    if not value:
        return None
    low = value.strip().lower()
    for m in MoneySubkind:
        if low == m.value:
            return m
    return _SUBKIND_BY_LABEL.get(low)


# Порядок важен: более специфичные / составные заголовки разбираются раньше,
# чтобы «Назначение платежа» не досталось короткому синониму «сумма платежа».
HEURISTIC_SYNONYMS: dict[str, tuple[str, ...]] = {
    "payment_purpose": ("назначение", "назначение платежа", "комментарий", "описание", "примечание", "purpose"),
    "doc_date": ("дата", "date", "дата операции", "дата документа", "дата платежа", "дата проводки"),
    "counterparty": ("контрагент", "плательщик", "получатель", "организация", "payer", "counterparty", "клиент"),
    "external_number": ("номер документа", "номер платежа", "номер поручения", "№ документа", "док. №", "number", " no", "номер"),
    "tax": ("ндс", "сумма ндс", "vat", " tax", "налог"),
    "subkind": ("вид платежа", "статья ддс", "статья", "категория", "тип платежа", "вид"),
    "direction_col": ("тип операции", "вид операции", "направление", "дебет/кредит", "приход/расход", "операция", "type"),
    "amount_debit": ("расход", "дебет", "списание", "списано", "debit", "outflow"),
    "amount_credit": ("приход", "кредит", "поступление", "поступило", "credit", "inflow"),
    "amount": ("сумма", "amount", "сумма платежа", "сумма операции", "сумма, руб", "сумма руб"),
}

CRITICAL_LABELS = {
    "amount": "с суммой платежа",
    "direction": "по которой видно приход/расход (знак суммы, тип операции или пара дебет/кредит)",
    "doc_date": "с датой документа",
    "counterparty": "с контрагентом",
    "tax": "с суммой НДС",
    "external_number": "с номером документа",
}


@dataclass
class PaymentColumnMapping:
    amount: str | None = None
    amount_debit: str | None = None
    amount_credit: str | None = None
    direction_col: str | None = None
    doc_date: str | None = None
    counterparty: str | None = None
    tax: str | None = None
    external_number: str | None = None
    subkind: str | None = None
    payment_purpose: str | None = None
    ai_used: bool = False
    note: str = ""

    def has_amount(self) -> bool:
        return bool(self.amount or (self.amount_debit and self.amount_credit))

    def has_direction_signal(self) -> bool:
        # знак суммы всегда даёт направление, если сумма есть одной колонкой;
        # пара дебет/кредит или явная колонка типа — тоже.
        return bool(self.amount or (self.amount_debit and self.amount_credit) or self.direction_col)

    def missing_critical(self) -> list[str]:
        miss: list[str] = []
        if not self.has_amount():
            miss.append("amount")
        if not self.has_direction_signal():
            miss.append("direction")
        for f in ("doc_date", "counterparty", "tax", "external_number"):
            if not getattr(self, f):
                miss.append(f)
        return miss

    def missing_optional(self) -> list[str]:
        return [f for f in OPTIONAL_FIELDS if not getattr(self, f)]

    def as_dict(self) -> dict:
        return {
            "amount": self.amount,
            "amount_debit": self.amount_debit,
            "amount_credit": self.amount_credit,
            "direction_col": self.direction_col,
            "doc_date": self.doc_date,
            "counterparty": self.counterparty,
            "tax": self.tax,
            "external_number": self.external_number,
            "subkind": self.subkind,
            "payment_purpose": self.payment_purpose,
        }


# --------------------------------------------------------------------------- #
# Чтение файла (та же механика, что в price_import)                           #
# --------------------------------------------------------------------------- #


def read_table(file: UploadFile) -> tuple[list[str], list[list[str]]]:
    raw = file.file.read()
    if not raw:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Пустой файл")

    name = (file.filename or "").lower()
    if name.endswith(".csv") or (not name.endswith((".xlsx", ".xls")) and _looks_like_csv(raw)):
        rows = _read_csv(raw)
    else:
        rows = _read_xlsx(raw)

    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="В файле нет строк с данными (нужны заголовок и хотя бы одна строка)",
        )
    headers = [str(c or "").strip() for c in rows[0]]
    if not any(headers):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="В первой строке файла нет заголовков")

    width = len(headers)
    data = [_pad(r, width) for r in rows[1 : 1 + MAX_DATA_ROWS]]
    return headers, data


def _looks_like_csv(raw: bytes) -> bool:
    head = raw[:4096]
    return b"," in head or b";" in head or b"\t" in head


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _read_csv(raw: bytes) -> list[list[str]]:
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.get_dialect("excel")
    reader = csv.reader(io.StringIO(text), dialect)
    return [[(c or "").strip() for c in row] for row in reader]


def _read_xlsx(raw: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(filename=io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Не удалось прочитать файл — ожидается .xlsx или .csv",
        ) from exc
    ws = wb.active
    out: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if v is None else str(v).strip() for v in row])
    wb.close()
    return out


def _pad(row: list[str], width: int) -> list[str]:
    row = [str(c or "").strip() for c in row]
    return (row + [""] * width)[:width]


# --------------------------------------------------------------------------- #
# Разметка колонок                                                            #
# --------------------------------------------------------------------------- #


def ai_enabled() -> bool:
    return bool(settings.anthropic_api_key)


def resolve_mapping(headers: list[str], sample: list[list[str]]) -> PaymentColumnMapping:
    if ai_enabled():
        try:
            mapping = _ai_mapping(headers, sample)
            mapping.ai_used = True
            return mapping
        except Exception:  # noqa: BLE001 — любой сбой ИИ = тихий откат
            pass
    mapping = _heuristic_mapping(headers)
    mapping.note = "разметка без ИИ (словарь синонимов)"
    return mapping


def _match_header(value: str | None, headers: list[str]) -> str | None:
    if not value:
        return None
    norm = value.strip().lower()
    for h in headers:
        if h.strip().lower() == norm:
            return h
    return None


def _heuristic_mapping(headers: list[str]) -> PaymentColumnMapping:
    mapping = PaymentColumnMapping()
    used: set[str] = set()
    for field_name, needles in HEURISTIC_SYNONYMS.items():
        for h in headers:
            if h in used:
                continue
            low = h.lower()
            # колонка «сумма» — не назначение платежа и не НДС
            if field_name == "amount" and ("назначен" in low or "ндс" in low or "vat" in low):
                continue
            if any(n.strip() in low for n in needles):
                setattr(mapping, field_name, h)
                used.add(h)
                break
    if mapping.amount and mapping.amount in (mapping.amount_debit, mapping.amount_credit):
        mapping.amount = None
    return mapping


_AI_TOOL = {
    "name": "map_payment_columns",
    "description": "Сопоставить заголовки таблицы платежей (выписка/реестр) с полями проводки.",
    "input_schema": {
        "type": "object",
        "properties": {
            "amount": {"type": "string", "description": "Заголовок колонки с суммой платежа одной колонкой (со знаком или без). Пусто, если суммы разнесены по дебету/кредиту."},
            "amount_debit": {"type": "string", "description": "Заголовок колонки «расход/дебет/списание», если сумма разнесена. Иначе пусто."},
            "amount_credit": {"type": "string", "description": "Заголовок колонки «приход/кредит/поступление», если сумма разнесена. Иначе пусто."},
            "direction_col": {"type": "string", "description": "Заголовок колонки с типом операции (приход/расход), если он текстом. Иначе пусто."},
            "doc_date": {"type": "string", "description": "Заголовок колонки с датой документа/операции."},
            "counterparty": {"type": "string", "description": "Заголовок колонки с наименованием контрагента (плательщик/получатель)."},
            "tax": {"type": "string", "description": "Заголовок колонки с суммой НДС."},
            "external_number": {"type": "string", "description": "Заголовок колонки с номером платёжного документа."},
            "subkind": {"type": "string", "description": "Заголовок колонки с видом/статьёй платежа, если есть. Иначе пусто."},
            "payment_purpose": {"type": "string", "description": "Заголовок колонки с назначением платежа. Иначе пусто."},
            "note": {"type": "string", "description": "Одно короткое предложение: что осталось неопознанным."},
        },
        "required": ["amount", "amount_debit", "amount_credit", "doc_date", "counterparty", "tax", "external_number"],
    },
}

_AI_SYSTEM = (
    "Ты сопоставляешь колонки загруженной таблицы платежей (банковская выписка "
    "или реестр) с полями проводки в системе учёта. Отвечай только вызовом "
    "инструмента map_payment_columns. Используй точные заголовки из присланного "
    "списка. Если подходящей колонки нет — оставляй поле пустым, не выдумывай."
)


def _ai_mapping(headers: list[str], sample: list[list[str]]) -> PaymentColumnMapping:
    from app.core.llm import anthropic_client

    preview = "\n".join(" | ".join(row) for row in sample[:SAMPLE_ROWS_FOR_AI])
    user = (
        f"Заголовки файла ({len(headers)}): {headers}\n\n"
        f"Первые строки данных:\n{preview}\n\n"
        "Сопоставь заголовки с полями: amount / amount_debit / amount_credit / "
        "direction_col / doc_date / counterparty / tax / external_number / "
        "subkind / payment_purpose."
    )
    client = anthropic_client(timeout=45.0, max_retries=2)
    response = client.messages.create(
        model=settings.ai_model,
        max_tokens=600,
        system=_AI_SYSTEM,
        messages=[{"role": "user", "content": user}],
        tools=[_AI_TOOL],
        tool_choice={"type": "tool", "name": "map_payment_columns"},
    )
    block = next((b for b in response.content if b.type == "tool_use"), None)
    if block is None:
        raise RuntimeError("no tool_use in map_payment_columns response")
    p = block.input or {}
    return PaymentColumnMapping(
        amount=_match_header(p.get("amount"), headers),
        amount_debit=_match_header(p.get("amount_debit"), headers),
        amount_credit=_match_header(p.get("amount_credit"), headers),
        direction_col=_match_header(p.get("direction_col"), headers),
        doc_date=_match_header(p.get("doc_date"), headers),
        counterparty=_match_header(p.get("counterparty"), headers),
        tax=_match_header(p.get("tax"), headers),
        external_number=_match_header(p.get("external_number"), headers),
        subkind=_match_header(p.get("subkind"), headers),
        payment_purpose=_match_header(p.get("payment_purpose"), headers),
        note=str(p.get("note") or "").strip(),
    )


# --------------------------------------------------------------------------- #
# Построение строк                                                            #
# --------------------------------------------------------------------------- #

_NUM_RE = re.compile(r"-?\d[\d\s]*(?:[.,]\d+)?")
_INCOME_WORDS = ("приход", "поступление", "кредит", "credit", "in", "зачисл")
_EXPENSE_WORDS = ("расход", "списание", "дебет", "debit", "out", "выплат")


def _to_number(value: str) -> float | None:
    if not value:
        return None
    m = _NUM_RE.search(value.replace("\xa0", " "))
    if not m:
        return None
    token = m.group(0).replace(" ", "").replace(",", ".")
    try:
        return float(token)
    except ValueError:
        return None


def _to_datetime(value: str) -> datetime | None:
    value = (value or "").strip()
    if not value:
        return None
    value = value.split("T")[0].split(" ")[0]
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(value, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


@dataclass
class BuiltRow:
    amount: float
    direction: MoneyDirection
    doc_date: datetime | None
    counterparty: str | None
    tax: float
    external_number: str | None
    subkind: MoneySubkind | None
    payment_purpose: str | None


@dataclass
class BuiltRows:
    items: list[BuiltRow] = field(default_factory=list)
    skipped: int = 0


def build_rows(headers: list[str], data: list[list[str]], mapping: PaymentColumnMapping) -> BuiltRows:
    idx = {h: i for i, h in enumerate(headers)}

    def cell(row: list[str], header: str | None) -> str:
        if not header or header not in idx:
            return ""
        i = idx[header]
        return row[i].strip() if i < len(row) else ""

    out = BuiltRows()
    for row in data:
        amount, direction = _row_amount_direction(row, mapping, cell)
        if amount is None:
            out.skipped += 1
            continue

        out.items.append(
            BuiltRow(
                amount=abs(amount),
                direction=direction,
                doc_date=_to_datetime(cell(row, mapping.doc_date)),
                counterparty=(cell(row, mapping.counterparty) or None),
                tax=_to_number(cell(row, mapping.tax)) or 0.0,
                external_number=(cell(row, mapping.external_number) or None),
                subkind=subkind_from_text(cell(row, mapping.subkind)),
                payment_purpose=(cell(row, mapping.payment_purpose) or None),
            )
        )
    return out


TEMPLATE_HEADERS = ["Дата", "Контрагент", "Назначение платежа", "Сумма", "НДС", "№ документа", "Вид"]


def generate_template() -> bytes:
    """.xlsx-шаблон таблицы платежей: заголовки + пара строк-примеров
    (приход со знаком «+», расход со знаком «−»)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Платежи"
    ws.append(TEMPLATE_HEADERS)
    ws.append(["01.09.2026", "ООО «Ромашка»", "оплата по счёту №12 от 25.08.2026", 150000, 25000, "125", "доход от продажи"])
    ws.append(["03.09.2026", "ИФНС №7", "НДС за 2 квартал 2026", -274000, 0, "126", "налоги и сборы"])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _row_amount_direction(row, mapping: PaymentColumnMapping, cell) -> tuple[float | None, MoneyDirection]:
    # 1) разнесённые дебет/кредит
    if mapping.amount_debit and mapping.amount_credit:
        debit = _to_number(cell(row, mapping.amount_debit))
        credit = _to_number(cell(row, mapping.amount_credit))
        if credit:
            return credit, MoneyDirection.INCOME
        if debit:
            return debit, MoneyDirection.EXPENSE
        return None, MoneyDirection.EXPENSE

    amount = _to_number(cell(row, mapping.amount))
    if amount is None:
        return None, MoneyDirection.EXPENSE

    # 2) явная колонка типа операции
    if mapping.direction_col:
        text = cell(row, mapping.direction_col).lower()
        if any(w in text for w in _INCOME_WORDS):
            return amount, MoneyDirection.INCOME
        if any(w in text for w in _EXPENSE_WORDS):
            return amount, MoneyDirection.EXPENSE

    # 3) знак суммы
    return amount, MoneyDirection.EXPENSE if amount < 0 else MoneyDirection.INCOME
