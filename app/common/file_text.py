"""Extract readable text from a stored FileAsset so Marina can use the spec."""

from __future__ import annotations

import logging
import os
import re

from app.common.files import FileAsset

log = logging.getLogger("app.common.file_text")

# DH96_Tech..., DH-96, dh_64B — без жёсткого \b после цифр: дальше часто идёт _
_MODEL = re.compile(r"DH[-_ ]?(\d+[A-Za-z]?)", re.IGNORECASE)
_MIN_NATIVE_CHARS = 80
_OCR_MAX_PAGES = 12
_OCR_DPI = 150
_IMAGE_TYPES = {"image/jpeg", "image/png", "image/webp", "image/tiff", "image/bmp"}
_IMAGE_SUFFIXES = (".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp")


def model_hint(filename: str) -> str | None:
    match = _MODEL.search(filename or "")
    if not match:
        return None
    return f"DH-{match.group(1).upper()}" if match.group(1)[0].isdigit() else f"DH-{match.group(1)}"


def extract_asset_text(asset: FileAsset, *, max_chars: int = 24000) -> dict:
    filename = asset.filename or ""
    hint = model_hint(filename)
    path = asset.path_on_disk
    source = "none"
    if not path or not os.path.isfile(path):
        return {
            "file_id": asset.id,
            "filename": filename,
            "purpose": asset.purpose.value,
            "model_hint": hint,
            "text": "",
            "source": source,
            "error": "Файл на диске не найден.",
        }
    content_type = (asset.content_type or "").lower()
    lower_name = filename.lower()
    text = ""

    if content_type.startswith("text/") or lower_name.endswith((".md", ".txt", ".csv", ".json")):
        text = open(path, encoding="utf-8", errors="replace").read()
        source = "text"
    elif content_type == "application/pdf" or lower_name.endswith(".pdf"):
        text = _pdf_text(path)
        source = "pdf"
        if _needs_ocr(text):
            ocr = _pdf_ocr(path)
            if len(ocr.strip()) > len(text.strip()):
                text = ocr
                source = "pdf+ocr"
    elif (
        content_type
        in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        }
        or lower_name.endswith((".docx", ".doc"))
    ):
        if lower_name.endswith(".docx"):
            text = _docx_text(path)
            source = "docx"
        else:
            text = "Старый формат .doc не читаю. Сохрани как .docx или PDF."
            source = "unsupported"
    elif content_type in _IMAGE_TYPES or lower_name.endswith(_IMAGE_SUFFIXES):
        text = _image_ocr(path)
        source = "image+ocr"
    else:
        text = ""

    if not text.strip():
        if hint:
            text = (
                f"Текст из файла не извлечён, но по имени это модель {hint}. "
                f"Опирайся на vault/search_company_vault по коду {hint} и на состав техкарты."
            )
        else:
            text = "Текст из файла не извлечён. Не утверждай, что спецификации нет — файл прикреплён."
        source = source if source != "none" else "fallback"
    return {
        "file_id": asset.id,
        "filename": filename,
        "purpose": asset.purpose.value,
        "model_hint": hint,
        "text": text[:max_chars],
        "truncated": len(text) > max_chars,
        "source": source,
    }


def _needs_ocr(text: str) -> bool:
    compact = re.sub(r"\s+", "", text or "")
    return len(compact) < _MIN_NATIVE_CHARS


def _pdf_text(path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return ""
    try:
        reader = PdfReader(path)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages).strip()
    except Exception:
        return ""


def _docx_text(path: str) -> str:
    try:
        from docx import Document
    except ImportError:
        return ""
    try:
        doc = Document(path)
        parts: list[str] = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _tesseract_langs() -> str:
    try:
        import pytesseract

        available = set(pytesseract.get_languages(config=""))
    except Exception:
        available = set()
    langs = [lang for lang in ("rus", "eng") if lang in available]
    return "+".join(langs) if langs else "eng"


def _ocr_image(image) -> str:
    try:
        import pytesseract
    except ImportError:
        return ""
    try:
        return pytesseract.image_to_string(image, lang=_tesseract_langs()) or ""
    except Exception as error:
        log.warning("OCR не сработал: %s", error)
        return ""


def _image_ocr(path: str) -> str:
    try:
        from PIL import Image
    except ImportError:
        return ""
    try:
        with Image.open(path) as image:
            return _ocr_image(image.convert("RGB")).strip()
    except Exception as error:
        log.warning("картинка для OCR не открылась: %s", error)
        return ""


def _pdf_ocr(path: str) -> str:
    try:
        import pymupdf
        from PIL import Image
    except ImportError:
        return ""
    try:
        doc = pymupdf.open(path)
    except Exception as error:
        log.warning("PDF для OCR не открылся: %s", error)
        return ""
    parts: list[str] = []
    try:
        page_count = min(len(doc), _OCR_MAX_PAGES)
        for index in range(page_count):
            page = doc[index]
            pix = page.get_pixmap(dpi=_OCR_DPI, alpha=False)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            chunk = _ocr_image(image).strip()
            if chunk:
                parts.append(chunk)
    finally:
        doc.close()
    return "\n\n".join(parts).strip()


def file_card(asset: FileAsset | None) -> dict | None:
    if asset is None:
        return None
    return {
        "id": asset.id,
        "filename": asset.filename,
        "content_type": asset.content_type,
        "purpose": asset.purpose.value,
        "model_hint": model_hint(asset.filename or ""),
    }
