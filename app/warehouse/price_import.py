"""Импорт прайс-листа поставщика таблицей (xlsx/csv) с ИИ-разметкой колонок.

Поток: читаем таблицу -> сопоставляем её заголовки с полями строки прайса
(`material`, `price`, `category`, `lead_time`, диапазоны партии). Сопоставление
делает Claude (forced tool). Если ключ ИИ не задан или вызов упал — откат на
словарь синонимов. Дальше строим строки прайса; недостающие необязательные поля
остаются пустыми, а раздел `service.import_price_list` ставит задачу «дозаполнить».
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from fastapi import HTTPException, UploadFile, status

from app.core.config import settings
from app.warehouse.models import MaterialCategory

MAX_DATA_ROWS = 1000
SAMPLE_ROWS_FOR_AI = 5

# Поля строки прайса, которые пытаемся найти в файле.
OPTIONAL_FIELDS = ("category", "lead_time")

# Откат без ИИ: подстрока заголовка (в нижнем регистре) -> поле.
HEURISTIC_SYNONYMS: dict[str, tuple[str, ...]] = {
    "material": ("наименование", "материал", "товар", "позиц", "номенклат", "продукц", "name", "item"),
    "price": ("цена", "стоимост", "прайс", "руб", "price", "cost", "сумма"),
    "category": ("категор", "групп", "раздел", "тип", "category"),
    "lead_time": ("срок", "поставк", "достав", "дн", "недел", "lead", "eta"),
}


@dataclass
class ColumnMapping:
    material: str | None = None
    price: str | None = None
    category: str | None = None
    lead_time: str | None = None
    # заголовки колонок «цена от N» / диапазоны партии
    qty_breaks: list[str] = field(default_factory=list)
    ai_used: bool = False
    note: str = ""

    def as_dict(self) -> dict:
        return {
            "material": self.material,
            "price": self.price,
            "category": self.category,
            "lead_time": self.lead_time,
            "qty_breaks": list(self.qty_breaks),
        }

    def missing_fields(self) -> list[str]:
        return [f for f in OPTIONAL_FIELDS if not getattr(self, f)]


# --------------------------------------------------------------------------- #
# Чтение файла                                                                #
# --------------------------------------------------------------------------- #


def read_table(file: UploadFile) -> tuple[list[str], list[list[str]]]:
    """(-> заголовки, строки данных как строки). Понимает .xlsx и .csv."""
    raw = file.file.read()
    if not raw:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Пустой файл")

    name = (file.filename or "").lower()
    if name.endswith(".csv") or (not name.endswith((".xlsx", ".xls")) and _looks_like_csv(raw)):
        rows = _read_csv(raw)
    else:
        rows = _read_xlsx(raw)

    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="В файле нет строк с данными (нужны заголовок и хотя бы одна строка)",
        )
    headers = [str(c or "").strip() for c in rows[0]]
    if not any(headers):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="В первой строке файла нет заголовков")

    width = len(headers)
    data = [_pad(r, width) for r in rows[1 : 1 + MAX_DATA_ROWS]]
    return headers, data


def _looks_like_csv(raw: bytes) -> bool:
    head = raw[:4096]
    return b"," in head or b";" in head or b"\t" in head


def _decode(raw: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _read_csv(raw: bytes) -> list[list[str]]:
    text = _decode(raw)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.get_dialect("excel")
    reader = csv.reader(io.StringIO(text), dialect)
    return [[(c or "").strip() for c in row] for row in reader]


def _read_xlsx(raw: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(filename=io.BytesIO(raw), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="Не удалось прочитать файл — ожидается .xlsx или .csv"
        ) from exc
    ws = wb.active
    out: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        out.append(["" if v is None else str(v).strip() for v in row])
    wb.close()
    return out


def _pad(row: list[str], width: int) -> list[str]:
    row = [str(c or "").strip() for c in row]
    return (row + [""] * width)[:width]


# --------------------------------------------------------------------------- #
# Сопоставление колонок                                                       #
# --------------------------------------------------------------------------- #


def resolve_mapping(headers: list[str], sample: list[list[str]]) -> ColumnMapping:
    """ИИ, если доступен; иначе — эвристика. Всегда возвращает ColumnMapping."""
    if ai_enabled():
        try:
            mapping = _ai_mapping(headers, sample)
            mapping.ai_used = True
            return mapping
        except Exception:  # noqa: BLE001 — любой сбой ИИ = тихий откат
            pass
    mapping = _heuristic_mapping(headers)
    mapping.note = "разметка без ИИ (словарь синонимов)"
    return mapping


def ai_enabled() -> bool:
    return bool(settings.anthropic_api_key)


def _match_header(value: str | None, headers: list[str]) -> str | None:
    """ИИ мог вернуть заголовок с точностью до регистра/пробелов — приводим к
    реальному заголовку файла, иначе None."""
    if not value:
        return None
    norm = value.strip().lower()
    for h in headers:
        if h.strip().lower() == norm:
            return h
    return None


def _heuristic_mapping(headers: list[str]) -> ColumnMapping:
    mapping = ColumnMapping()
    # Сначала снимаем колонки-диапазоны («цена от 100», «опт от 1000»), чтобы
    # синоним «цена» не забрал их себе как обычную колонку цены.
    breaks = [h for h in headers if _qty_break_threshold(h) is not None]
    mapping.qty_breaks = breaks
    used: set[str] = set(breaks)
    for field_name, needles in HEURISTIC_SYNONYMS.items():
        for h in headers:
            if h in used:
                continue
            low = h.lower()
            if any(n in low for n in needles):
                setattr(mapping, field_name, h)
                used.add(h)
                break
    return mapping


_QTY_BREAK_RE = re.compile(r"(?:от|>=?|более|свыше)?\s*(\d[\d\s]*)\s*(?:\+|шт|штук|ед|уп)?", re.IGNORECASE)


def _qty_break_threshold(header: str) -> float | None:
    low = header.lower()
    if not any(k in low for k in ("цена", "опт", "стоим", "price", "руб")):
        return None
    if not any(k in low for k in ("от", "+", ">", "более", "свыше", "парти", "объ")):
        return None
    m = _QTY_BREAK_RE.search(header)
    if not m:
        return None
    digits = m.group(1).replace(" ", "")
    return float(digits) if digits.isdigit() else None


_AI_TOOL = {
    "name": "map_columns",
    "description": "Сопоставить заголовки прайс-листа поставщика с полями системы.",
    "input_schema": {
        "type": "object",
        "properties": {
            "material": {"type": "string", "description": "Точный заголовок колонки с названием материала/товара. Пусто, если такой колонки нет."},
            "price": {"type": "string", "description": "Точный заголовок колонки с ценой за единицу. Пусто, если такой колонки нет."},
            "category": {"type": "string", "description": "Заголовок колонки с категорией/группой. Пусто, если нет."},
            "lead_time": {"type": "string", "description": "Заголовок колонки со сроком поставки. Пусто, если нет."},
            "qty_breaks": {
                "type": "array",
                "items": {"type": "string"},
                "description": "Заголовки колонок с ценами по объёму партии (напр. «цена от 100»), если такие есть.",
            },
            "note": {"type": "string", "description": "Одно короткое предложение: что осталось неопознанным."},
        },
        "required": ["material", "price", "qty_breaks"],
    },
}

_AI_SYSTEM = (
    "Ты сопоставляешь колонки загруженного прайс-листа поставщика с полями "
    "системы учёта. Отвечай только вызовом инструмента map_columns. Используй "
    "точные заголовки из присланного списка. Если подходящей колонки нет — "
    "оставляй поле пустым, не выдумывай."
)


def _ai_mapping(headers: list[str], sample: list[list[str]]) -> ColumnMapping:
    from app.core.llm import anthropic_client

    preview = "\n".join(" | ".join(row) for row in sample[:SAMPLE_ROWS_FOR_AI])
    user = (
        f"Заголовки файла ({len(headers)}): {headers}\n\n"
        f"Первые строки данных:\n{preview}\n\n"
        "Сопоставь заголовки с полями: material, price, category, lead_time, "
        "qty_breaks."
    )
    client = anthropic_client(timeout=45.0, max_retries=2)
    response = client.messages.create(
        model=settings.ai_model,
        max_tokens=512,
        system=_AI_SYSTEM,
        messages=[{"role": "user", "content": user}],
        tools=[_AI_TOOL],
        tool_choice={"type": "tool", "name": "map_columns"},
    )
    block = next((b for b in response.content if b.type == "tool_use"), None)
    if block is None:
        raise RuntimeError("no tool_use in map_columns response")
    payload = block.input or {}

    raw_breaks = payload.get("qty_breaks") or []
    qty_breaks = [h for h in (_match_header(b, headers) for b in raw_breaks) if h]
    return ColumnMapping(
        material=_match_header(payload.get("material"), headers),
        price=_match_header(payload.get("price"), headers),
        category=_match_header(payload.get("category"), headers),
        lead_time=_match_header(payload.get("lead_time"), headers),
        qty_breaks=qty_breaks,
        note=str(payload.get("note") or "").strip(),
    )


# --------------------------------------------------------------------------- #
# Построение строк прайса                                                     #
# --------------------------------------------------------------------------- #

_NUM_RE = re.compile(r"-?\d[\d\s]*(?:[.,]\d+)?")


def _to_number(value: str) -> float | None:
    if not value:
        return None
    m = _NUM_RE.search(value.replace("\xa0", " "))
    if not m:
        return None
    token = m.group(0).replace(" ", "").replace(",", ".")
    try:
        return float(token)
    except ValueError:
        return None


@dataclass
class BuiltRows:
    items: list[dict]
    skipped: int


def build_rows(headers: list[str], data: list[list[str]], mapping: ColumnMapping) -> BuiltRows:
    """mapping.material и mapping.price должны быть заданы (проверяется в роутере)."""
    idx = {h: i for i, h in enumerate(headers)}
    m_i = idx[mapping.material]  # type: ignore[index]
    p_i = idx.get(mapping.price) if mapping.price else None
    c_i = idx.get(mapping.category) if mapping.category else None
    l_i = idx.get(mapping.lead_time) if mapping.lead_time else None
    breaks = sorted(
        ((h, _qty_break_threshold(h) or 0.0) for h in mapping.qty_breaks if h in idx),
        key=lambda x: x[1],
    )

    items: list[dict] = []
    skipped = 0
    for row in data:
        material = (row[m_i] if m_i < len(row) else "").strip()
        if not material:
            skipped += 1
            continue

        tiers: list[dict] = []
        if breaks:
            for pos, (h, threshold) in enumerate(breaks):
                price = _to_number(row[idx[h]] if idx[h] < len(row) else "")
                if price is None:
                    continue
                nxt = breaks[pos + 1][1] if pos + 1 < len(breaks) else None
                tiers.append({"min_qty": threshold, "max_qty": nxt, "price": price})
        else:
            price = _to_number(row[p_i] if p_i is not None and p_i < len(row) else "")
            if price is not None:
                tiers.append({"min_qty": 0, "max_qty": None, "price": price})

        if not tiers:
            skipped += 1
            continue

        items.append(
            {
                "material": material[:255],
                "category": ((row[c_i].strip() or None) if c_i is not None and c_i < len(row) else None),
                "lead_time": ((row[l_i].strip() or None) if l_i is not None and l_i < len(row) else None),
                "tiers": tiers,
                "round": None,
            }
        )
    return BuiltRows(items=items, skipped=skipped)


# --------------------------------------------------------------------------- #
# ИИ-помощь по недостающим полям (задача 0011-h)                              #
# --------------------------------------------------------------------------- #

CATEGORY_CHOICES = [c.value for c in MaterialCategory if c is not MaterialCategory.NONE]

_CATEGORY_TOOL = {
    "name": "assign_categories",
    "description": "Проставить категорию складского справочника каждой позиции по её названию.",
    "input_schema": {
        "type": "object",
        "properties": {
            "assignments": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "id": {"type": "integer"},
                        "category": {
                            "type": "string",
                            "description": "Строго одно из значений справочника; пусто, если не определяется.",
                        },
                    },
                    "required": ["id", "category"],
                },
            }
        },
        "required": ["assignments"],
    },
}


def ai_assign_categories(pairs: list[tuple[int, str]]) -> dict[int, str]:
    """{id строки прайса -> категория из справочника}. Только уверенные назначения."""
    if not pairs:
        return {}
    from app.core.llm import anthropic_client

    listing = "\n".join(f"{pid}. {material}" for pid, material in pairs)
    user = (
        "Справочник категорий склада:\n- " + "\n- ".join(CATEGORY_CHOICES) + "\n\n"
        "Позиции прайса (id. название):\n" + listing + "\n\n"
        "Для каждой позиции выбери ближайшую категорию из справочника. Если "
        "категория не определяется однозначно — оставь пустую строку."
    )
    client = anthropic_client(timeout=45.0, max_retries=2)
    response = client.messages.create(
        model=settings.ai_model,
        max_tokens=1024,
        system="Ты классифицируешь строительные материалы по складскому справочнику. Отвечай только вызовом assign_categories.",
        messages=[{"role": "user", "content": user}],
        tools=[_CATEGORY_TOOL],
        tool_choice={"type": "tool", "name": "assign_categories"},
    )
    block = next((b for b in response.content if b.type == "tool_use"), None)
    if block is None:
        raise RuntimeError("no tool_use in assign_categories response")

    allowed = set(CATEGORY_CHOICES)
    ids = {pid for pid, _ in pairs}
    out: dict[int, str] = {}
    for row in (block.input or {}).get("assignments") or []:
        try:
            pid = int(row.get("id"))
        except (TypeError, ValueError):
            continue
        cat = str(row.get("category") or "").strip()
        if pid in ids and cat in allowed:
            out[pid] = cat
    return out


def ai_lead_time_message(supplier_name: str, materials: list[str]) -> str:
    """Короткое вежливое сообщение поставщику с просьбой указать сроки поставки."""
    from app.core.llm import anthropic_client

    listing = "\n".join(f"{i}. {m}" for i, m in enumerate(materials, start=1))
    client = anthropic_client(timeout=45.0, max_retries=2)
    response = client.messages.create(
        model=settings.ai_model,
        max_tokens=600,
        system=(
            "Ты менеджер по снабжению. Напиши короткое вежливое сообщение поставщику "
            "в мессенджере: попроси указать срок поставки по каждой позиции из списка. "
            "Без формального шапки-подписи, простой деловой тон, по-русски. Верни только текст сообщения."
        ),
        messages=[
            {
                "role": "user",
                "content": (
                    f"Поставщик: {supplier_name}. Нужны сроки поставки по позициям:\n{listing}\n\n"
                    "Вставь этот список в сообщение и попроси напротив каждой позиции указать срок."
                ),
            }
        ],
    )
    text = "".join(getattr(b, "text", "") for b in response.content if b.type == "text").strip()
    if not text:
        raise RuntimeError("empty lead-time message")
    return text
